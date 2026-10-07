from docx.document import Document as WordDocument
from docx.shared import Inches
from httpx import AsyncClient

from tests.documents import DOCX, add_word_xml, read_content, word
from tests.fakes import FakeBoundary, text_file


def quarterly_report(document: WordDocument) -> None:
    document.add_heading("Quarterly report", 0)
    document.add_heading("Results", 1)
    document.add_paragraph("Revenue grew by 12%.")
    document.add_paragraph("New customers", style="List Bullet")
    document.add_paragraph("Fewer returns", style="List Number")
    table = document.add_table(rows=2, cols=2)
    for row, cells in enumerate([("Region", "Revenue"), ("North", "1200")]):
        for column, text in enumerate(cells):
            table.cell(row, column).text = text
    document.add_heading("Outlook", 2)
    document.add_paragraph("Stable.")


async def test_a_word_document_comes_as_headings_paragraphs_and_tables(
    client: AsyncClient, boundary: FakeBoundary
) -> None:
    content = word(quarterly_report)
    boundary.drive.add(text_file("report", "Report.docx", content=content, mime=DOCX))

    response = await read_content(client, "report")

    assert response.status_code == 200, response.text
    # Its headings as Markdown's, its tables as rows of cells, and nothing else of what it holds
    assert response.json() == {
        "id": "report",
        "size": len(content),
        "truncated": False,
        "untrusted": {
            "name": "Report.docx",
            "mime": DOCX,
            "content": "# Quarterly report\n"
            "\n"
            "# Results\n"
            "Revenue grew by 12%.\n"
            "- New customers\n"
            "- Fewer returns\n"
            "\n"
            "Region\tRevenue\n"
            "North\t1200\n"
            "\n"
            "## Outlook\n"
            "Stable.",
        },
    }


def marked_up(document: WordDocument) -> None:
    # Tab stops sit with the paragraph's properties, apart from the tabs of its text
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(1))
    paragraph.add_run("Name\tValue").add_break()
    paragraph.add_run("next line")
    add_word_xml(
        document,
        '<w:p><w:r><w:t xml:space="preserve">Kept </w:t></w:r>'
        '<w:del w:id="1" w:author="Alice"><w:r><w:delText>removed </w:delText></w:r></w:del>'
        '<w:ins w:id="2" w:author="Alice"><w:r><w:t>added</w:t></w:r></w:ins></w:p>',
    )
    add_word_xml(
        document,
        '<w:p><w:r><w:t xml:space="preserve">Page </w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>7</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>',
    )
    # A text box, as Word writes it twice: for the readers that know its shapes, and for others
    add_word_xml(
        document,
        '<w:p><w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:txbx>'
        "<w:txbxContent><w:p><w:r><w:t>Boxed words</w:t></w:r></w:p></w:txbxContent>"
        "</wps:txbx></wps:wsp></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox>"
        "<w:txbxContent><w:p><w:r><w:t>Boxed words</w:t></w:r></w:p></w:txbxContent>"
        "</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
        "<w:r><w:t>Around the box.</w:t></w:r></w:p>",
    )


async def test_the_text_a_reader_sees_comes_once(
    client: AsyncClient, boundary: FakeBoundary
) -> None:
    boundary.drive.add(text_file("doc", "Doc.docx", content=word(marked_up), mime=DOCX))

    response = await read_content(client, "doc")

    assert response.status_code == 200, response.text
    # Neither deleted text, nor field codes, nor the copy of a text box for older readers
    assert response.json()["untrusted"]["content"] == (
        "Name\tValue\nnext line\nKept added\nPage 7\nBoxed words\nAround the box."
    )


async def test_a_carriage_return_in_a_document_ends_its_line(
    client: AsyncClient, boundary: FakeBoundary
) -> None:
    def returned(document: WordDocument) -> None:
        add_word_xml(document, "<w:p><w:r><w:t>Shown&#13;Written over</w:t></w:r></w:p>")

    boundary.drive.add(text_file("doc", "Doc.docx", content=word(returned), mime=DOCX))

    response = await read_content(client, "doc")

    assert response.json()["untrusted"]["content"] == "Shown\nWritten over"


async def test_what_a_reader_does_not_see_is_left_out(
    client: AsyncClient, boundary: FakeBoundary
) -> None:
    content = word(lambda document: document.add_paragraph("A‮B​C\x85D\x9bE"))
    boundary.drive.add(text_file("doc", "Doc.docx", content=content, mime=DOCX))

    response = await read_content(client, "doc")

    # Neither the format characters that reorder or hide text, nor control characters
    assert response.json()["untrusted"]["content"] == "ABCDE"


async def test_a_long_document_is_cut_at_max_bytes(
    client: AsyncClient, boundary: FakeBoundary
) -> None:
    content = word(lambda document: document.add_paragraph("é" * 10))
    boundary.drive.add(text_file("doc", "Doc.docx", content=content, mime=DOCX))

    answer = (await read_content(client, "doc", max_bytes=5)).json()

    # Five bytes of its text hold two é and half of the third, which is left out
    assert answer["untrusted"]["content"] == "éé"
    assert answer["truncated"] is True
    assert answer["size"] == len(content)


async def test_a_document_longer_than_what_is_asked_is_cut(
    client: AsyncClient, boundary: FakeBoundary
) -> None:
    def long(document: WordDocument) -> None:
        for number in range(2_000):
            document.add_paragraph(f"Paragraph {number} of a long document.")

    boundary.drive.add(text_file("doc", "Doc.docx", content=word(long), mime=DOCX))

    answer = (await read_content(client, "doc", max_bytes=1_000)).json()

    assert answer["truncated"] is True
    assert answer["untrusted"]["content"].startswith("Paragraph 0 of a long document.\n")
    assert 900 < len(answer["untrusted"]["content"].encode()) <= 1_000


async def test_a_damaged_word_document_is_not_extractable(
    client: AsyncClient, boundary: FakeBoundary
) -> None:
    content = word(lambda document: document.add_paragraph("Confidential plans"))
    boundary.drive.add(
        text_file("cut", "Cut.docx", content=content[: len(content) // 2], mime=DOCX),
        text_file("other", "Other.docx", content=b"Confidential plans", mime=DOCX),
    )

    for file_id in ("cut", "other"):
        response = await read_content(client, file_id)

        assert response.status_code == 415, response.text
        assert response.json()["code"] == "content_not_extractable"
        # What the file holds stays out of the problem
        assert "Confidential" not in response.text
